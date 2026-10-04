"""
File Parser — 뒤죽박죭 폴더의 다양한 파일 형식을 텍스트로 변환
지원: .hwp, .hwpx, .docx, .xlsx, .pdf, .txt, .csv, .md
"""

import os
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Optional
from dataclasses import dataclass, field
from datetime import datetime

from docx import Document as DocxDocument
from openpyxl import load_workbook
from PyPDF2 import PdfReader


@dataclass
class ParsedFile:
    """파싱된 파일 정보"""
    filename: str
    filepath: str
    extension: str
    content: str                          # 추출된 텍스트
    metadata: dict = field(default_factory=dict)  # 작성자, 날짜 등
    size_bytes: int = 0
    page_count: Optional[int] = None
    error: Optional[str] = None

    @property
    def content_preview(self) -> str:
        """앞 500자 미리보기"""
        return self.content[:500] + "..." if len(self.content) > 500 else self.content


class FileParser:
    """다양한 형식의 파일을 텍스트로 변환하는 파서"""

    SUPPORTED_EXTENSIONS = {".hwp", ".hwpx", ".docx", ".xlsx", ".pdf", ".txt", ".csv", ".md"}

    def parse_file(self, filepath: str) -> ParsedFile:
        """단일 파일 파싱"""
        path = Path(filepath)
        ext = path.suffix.lower()
        size = path.stat().st_size if path.exists() else 0

        base = ParsedFile(
            filename=path.name,
            filepath=str(path),
            extension=ext,
            content="",
            size_bytes=size,
        )

        try:
            if ext == ".hwpx":
                return self._parse_hwpx(path, base)
            elif ext == ".hwp":
                return self._parse_hwp(path, base)
            elif ext == ".docx":
                return self._parse_docx(path, base)
            elif ext == ".xlsx":
                return self._parse_xlsx(path, base)
            elif ext == ".pdf":
                return self._parse_pdf(path, base)
            elif ext in (".txt", ".csv", ".md"):
                return self._parse_text(path, base)
            else:
                base.error = f"미지원 형식: {ext}"
                return base
        except Exception as e:
            base.error = f"파싱 오류: {str(e)}"
            return base

    def parse_folder(self, folder_path: str) -> list[ParsedFile]:
        """폴더 전체를 재귀적으로 파싱"""
        results = []
        folder = Path(folder_path)

        if not folder.exists():
            return results

        for file_path in sorted(folder.rglob("*")):
            if file_path.is_file() and file_path.suffix.lower() in self.SUPPORTED_EXTENSIONS:
                parsed = self.parse_file(str(file_path))
                # 상대경로 저장 (폴더 구조 자체가 업무 정보)
                parsed.metadata["relative_path"] = str(file_path.relative_to(folder))
                parsed.metadata["parent_folder"] = str(file_path.parent.relative_to(folder))
                results.append(parsed)

        return results

    def _parse_hwpx(self, path: Path, base: ParsedFile) -> ParsedFile:
        """HWPX 파싱 — 표(tbl) 구조를 보존하여 텍스트 추출"""
        texts = []

        with zipfile.ZipFile(str(path), "r") as zf:
            xml_files = sorted(
                [n for n in zf.namelist() if n.startswith("Contents/sec") and n.endswith(".xml")]
            )
            if not xml_files:
                xml_files = [n for n in zf.namelist() if n.endswith(".xml")]

            for xml_name in xml_files:
                try:
                    xml_data = zf.read(xml_name)
                    root = ET.fromstring(xml_data)
                    texts.extend(self._hwpx_walk(root))
                except ET.ParseError:
                    continue

            if "META-INF/manifest.xml" in zf.namelist():
                base.metadata["format"] = "HWPX (OWPML)"

        base.content = "\n".join(texts)
        base.metadata["parse_method"] = "hwpx_xml_table_aware"
        return base

    def _hwpx_walk(self, node) -> list[str]:
        """HWPX XML을 재귀 순회하며 표 구조를 마크다운 테이블로 변환"""
        results = []
        tag = node.tag.split("}")[-1] if "}" in node.tag else node.tag

        # 표(tbl) 처리 — 행/열 구조를 마크다운 테이블로 보존
        if tag == "tbl":
            table_rows = []
            for child in node:
                child_tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                if child_tag == "tr":
                    cells = []
                    for cell_node in child:
                        cell_tag = cell_node.tag.split("}")[-1] if "}" in cell_node.tag else cell_node.tag
                        if cell_tag == "tc":
                            cell_text = self._hwpx_extract_text(cell_node).strip()
                            cells.append(cell_text if cell_text else " ")
                    if cells:
                        table_rows.append(cells)

            if table_rows:
                results.append("\n[표]")
                # 헤더 행
                results.append("| " + " | ".join(table_rows[0]) + " |")
                results.append("| " + " | ".join(["---"] * len(table_rows[0])) + " |")
                # 데이터 행
                for row in table_rows[1:]:
                    # 열 수 맞추기
                    while len(row) < len(table_rows[0]):
                        row.append(" ")
                    results.append("| " + " | ".join(row[:len(table_rows[0])]) + " |")
                results.append("")
            return results

        # 문단(p) 처리
        if tag == "p":
            text = self._hwpx_extract_text(node).strip()
            if text:
                results.append(text)
            return results

        # 재귀 순회
        for child in node:
            results.extend(self._hwpx_walk(child))

        return results

    def _hwpx_extract_text(self, node) -> str:
        """HWPX 노드에서 텍스트만 추출 (재귀)"""
        parts = []
        tag = node.tag.split("}")[-1] if "}" in node.tag else node.tag

        if tag == "t" and node.text:
            parts.append(node.text)

        for child in node:
            parts.append(self._hwpx_extract_text(child))

        if node.tail and node.tail.strip():
            parts.append(node.tail.strip())

        return "".join(parts)

    def _parse_hwp(self, path: Path, base: ParsedFile) -> ParsedFile:
        """HWP 파싱 — LibreOffice로 텍스트 변환"""
        with tempfile.TemporaryDirectory() as tmp_dir:
            try:
                result = subprocess.run(
                    [
                        "libreoffice",
                        "--headless",
                        "--convert-to", "txt:Text (encoded):UTF8",
                        "--outdir", tmp_dir,
                        str(path),
                    ],
                    capture_output=True,
                    text=True,
                    timeout=30,
                )

                # 변환된 txt 파일 찾기
                txt_files = list(Path(tmp_dir).glob("*.txt"))
                if txt_files:
                    content = txt_files[0].read_text(encoding="utf-8-sig")  # BOM 제거
                    base.content = content.strip()
                    base.metadata["encoding"] = "utf-8"
                else:
                    base.error = f"LibreOffice 변환 실패: {result.stderr[:200]}"

            except subprocess.TimeoutExpired:
                base.error = "HWP 변환 시간 초과 (30초)"
            except FileNotFoundError:
                base.error = "LibreOffice가 설치되어 있지 않습니다"

        base.metadata["parse_method"] = "libreoffice"
        return base

    def _parse_docx(self, path: Path, base: ParsedFile) -> ParsedFile:
        """Word 문서 파싱"""
        doc = DocxDocument(str(path))

        # 메타데이터 추출
        props = doc.core_properties
        base.metadata.update({
            "author": props.author or "",
            "created": str(props.created) if props.created else "",
            "modified": str(props.modified) if props.modified else "",
            "title": props.title or "",
            "subject": props.subject or "",
        })

        # 본문 텍스트
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # 표 텍스트
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    table_texts.append(row_text)

        all_text = "\n".join(paragraphs)
        if table_texts:
            all_text += "\n\n[표]\n" + "\n".join(table_texts)

        base.content = all_text
        return base

    def _parse_xlsx(self, path: Path, base: ParsedFile) -> ParsedFile:
        """엑셀 파일 파싱"""
        wb = load_workbook(str(path), read_only=True, data_only=True)

        sheets_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(max_row=200, values_only=True):  # 최대 200행
                row_vals = [str(cell) if cell is not None else "" for cell in row]
                if any(v.strip() for v in row_vals):
                    rows.append(" | ".join(row_vals))

            if rows:
                sheets_text.append(f"[시트: {sheet_name}]\n" + "\n".join(rows))

        wb.close()
        base.content = "\n\n".join(sheets_text)
        base.metadata["sheet_count"] = len(wb.sheetnames)
        return base

    def _parse_pdf(self, path: Path, base: ParsedFile) -> ParsedFile:
        """PDF 파싱"""
        reader = PdfReader(str(path))
        base.page_count = len(reader.pages)

        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                pages_text.append(f"[페이지 {i+1}]\n{text.strip()}")

        base.content = "\n\n".join(pages_text)

        # 메타데이터
        if reader.metadata:
            base.metadata.update({
                "author": reader.metadata.get("/Author", ""),
                "title": reader.metadata.get("/Title", ""),
                "created": str(reader.metadata.get("/CreationDate", "")),
            })

        return base

    def _parse_text(self, path: Path, base: ParsedFile) -> ParsedFile:
        """텍스트 파일 파싱"""
        encodings = ["utf-8", "euc-kr", "cp949"]
        for enc in encodings:
            try:
                base.content = path.read_text(encoding=enc)
                base.metadata["encoding"] = enc
                return base
            except UnicodeDecodeError:
                continue

        base.error = "인코딩 인식 실패"
        return base


def get_folder_summary(parsed_files: list[ParsedFile]) -> dict:
    """파싱 결과의 요약 통계"""
    total = len(parsed_files)
    by_ext = {}
    errors = []
    total_chars = 0

    for pf in parsed_files:
        by_ext[pf.extension] = by_ext.get(pf.extension, 0) + 1
        total_chars += len(pf.content)
        if pf.error:
            errors.append({"file": pf.filename, "error": pf.error})

    # 폴더 구조 추출
    folders = set()
    for pf in parsed_files:
        parent = pf.metadata.get("parent_folder", "")
        if parent and parent != ".":
            folders.add(parent)

    return {
        "total_files": total,
        "by_extension": by_ext,
        "total_characters": total_chars,
        "folder_structure": sorted(folders),
        "errors": errors,
    }
