"""
Exporter — 인수인계서를 DOCX / PDF로 내보내기
"""

import re
import subprocess
import tempfile
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class HandoverExporter:
    """마크다운 인수인계서를 DOCX/PDF로 변환"""

    def to_docx(self, markdown: str, output_path: str) -> str:
        """Markdown → DOCX 변환"""
        doc = Document()

        # 기본 스타일 설정
        style = doc.styles["Normal"]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)

        lines = markdown.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i]

            # 제목
            if line.startswith("# "):
                p = doc.add_heading(line[2:].strip(), level=1)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)

            # 마크다운 테이블
            elif line.strip().startswith("|") and "|" in line:
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    # 구분선(---|---) 건너뛰기
                    if not re.match(r"^\|[\s\-|]+\|$", lines[i].strip()):
                        table_lines.append(lines[i])
                    i += 1
                i -= 1  # 외부 루프에서 +1 하므로

                if table_lines:
                    self._add_table(doc, table_lines)

            # 구분선
            elif line.strip() == "---":
                p = doc.add_paragraph()
                p.add_run("─" * 60).font.color.rgb = RGBColor(200, 200, 200)

            # 목록
            elif line.strip().startswith("- "):
                text = line.strip()[2:]
                text = self._clean_bold(text)
                p = doc.add_paragraph(text, style="List Bullet")

            # 볼드 텍스트 라인
            elif line.strip().startswith("**") and line.strip().endswith("**"):
                text = line.strip()[2:-2]
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True

            # 일반 텍스트
            elif line.strip():
                text = self._clean_bold(line.strip())
                doc.add_paragraph(text)

            # 빈 줄
            else:
                pass

            i += 1

        doc.save(output_path)
        return output_path

    def to_pdf(self, markdown: str, output_path: str) -> str:
        """Markdown → PDF (DOCX 경유 LibreOffice 변환)"""
        with tempfile.TemporaryDirectory() as tmp_dir:
            docx_path = str(Path(tmp_dir) / "handover.docx")
            self.to_docx(markdown, docx_path)

            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "pdf",
                 "--outdir", str(Path(output_path).parent), docx_path],
                capture_output=True, timeout=30,
            )

            # LibreOffice 출력 파일명 맞추기
            generated_pdf = Path(output_path).parent / "handover.pdf"
            if generated_pdf.exists() and str(generated_pdf) != output_path:
                generated_pdf.rename(output_path)

        return output_path

    def _add_table(self, doc, table_lines):
        """마크다운 테이블 라인을 DOCX 테이블로 변환"""
        rows_data = []
        for line in table_lines:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            rows_data.append(cells)

        if not rows_data:
            return

        cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=cols)
        table.style = "Light Grid Accent 1"

        for r_idx, row in enumerate(rows_data):
            for c_idx, cell_text in enumerate(row):
                if c_idx < cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_text
                    for paragraph in cell.paragraphs:
                        paragraph.style.font.size = Pt(10)
                        if r_idx == 0:
                            for run in paragraph.runs:
                                run.bold = True

        doc.add_paragraph()  # 테이블 후 간격

    def _clean_bold(self, text: str) -> str:
        """마크다운 볼드(**) 제거"""
        return re.sub(r"\*\*(.+?)\*\*", r"\1", text)
